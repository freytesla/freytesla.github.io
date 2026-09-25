from pathlib import Path

from docx import Document
from docx.oxml.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.shared import Pt


source = Path(r"D:\Desktop\frey.github.io\数据结构课程设计选题说明表_迷宫寻路_已填写.docx")
output = Path(r"D:\Desktop\frey.github.io\数据结构课程设计选题说明表_迷宫寻路_10月26日汇报版.docx")

doc = Document(source)
cell = doc.tables[0].rows[5].cells[0]
cell.text = ""

content = [
    ("三、项目实施步骤及进度安排", True),
    ("1. 需求分析与总体设计（2026.09.14—2026.09.20）：", True),
    ("明确输入输出格式、功能菜单和异常规则，完成数据结构、程序模块及测试方案设计。", False),
    ("2. 核心功能编码（2026.09.21—2026.10.11）：", True),
    ("完成迷宫输入、显示、DFS 可行路径、BFS 最短路径和路径回溯模块。", False),
    ("3. 测试与优化（2026.10.12—2026.10.18）：", True),
    ("测试不同规模及边界情况，修复错误，完善随机迷宫、算法对比和可视化功能。", False),
    ("4. 文档与答辩准备（2026.10.19—2026.10.25）：", True),
    ("整理测试数据和运行截图，完成课程设计报告、程序使用说明与答辩 PPT，并进行汇报演练。", False),
    ("5. 成果汇报（2026.10.26）：", True),
    ("上台展示系统功能、算法设计、测试结果，并完成课程设计答辩。", False),
]

paragraphs = [cell.paragraphs[0]]
for _ in range(len(content) - 1):
    paragraphs.append(cell.add_paragraph())

for paragraph, (text, bold) in zip(paragraphs, content):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)

doc.save(output)
print(output)
